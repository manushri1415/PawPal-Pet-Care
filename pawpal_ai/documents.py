"""Document validation, text extraction, and prompt-injection scanning.

The first guardrail in the pipeline. This module treats uploaded files as
*untrusted data*: it validates them, extracts text without executing anything,
and scans for prompt-injection patterns so downstream prompts can wrap the text
in explicit "this is data, not instructions" delimiters. Every failure path
returns a structured result instead of raising into the UI.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from pawpal_ai.logging_setup import log_event

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}
MAX_BYTES = 5 * 1024 * 1024  # 5 MB cap
MIN_CHARS = 10  # anything shorter is treated as effectively empty

# Patterns that suggest a document is trying to hijack the model. We only *flag*
# these — the real defense is the untrusted-data framing + mandatory human
# approval before anything is saved.
_INJECTION_PATTERNS = [
    r"ignore (all |the )?(previous|prior|above) (instructions|prompts?)",
    r"disregard (the |all )?(previous|prior|system) ",
    r"\bsystem prompt\b",
    r"^\s*(system|assistant)\s*:",
    r"you are now",
    r"forget (everything|all previous)",
    r"</?(instructions?|system)>",
]
_INJECTION_RE = re.compile("|".join(_INJECTION_PATTERNS), re.IGNORECASE | re.MULTILINE)


@dataclass
class DocumentResult:
    """Outcome of ingesting one document."""

    ok: bool
    text: str = ""
    doc_type: str = ""
    filename: str = ""
    char_count: int = 0
    injection_flagged: bool = False
    error: Optional[str] = None
    injection_spans: list[str] = field(default_factory=list)


def scan_for_injection(text: str) -> list[str]:
    """Return the (truncated) matched spans of any injection-looking content."""
    spans = [m.group(0).strip()[:80] for m in _INJECTION_RE.finditer(text)]
    return spans


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts = []
    for page in reader.pages:
        parts.append(page.extract_text() or "")
    return "\n\n".join(parts)


def _extract_docx(data: bytes) -> str:
    import docx  # python-docx

    document = docx.Document(io.BytesIO(data))
    paras = [p.text for p in document.paragraphs]
    # Flatten simple tables (vet records are often tabular).
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            paras.append(": ".join([c for c in cells if c]))
    return "\n".join(paras)


def _extract_txt(data: bytes) -> str:
    for enc in ("utf-8", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ValueError("could not decode text file")


def ingest_bytes(data: bytes, filename: str) -> DocumentResult:
    """Validate + extract text from raw bytes. Never raises for expected
    failures (empty/oversize/unsupported/corrupt) — returns ``ok=False``."""
    ext = Path(filename).suffix.lower()

    if ext not in SUPPORTED_EXTENSIONS:
        log_event("document_rejected", reason="unsupported_type", ext=ext)
        return DocumentResult(ok=False, filename=filename, error=f"Unsupported file type '{ext}'. Use PDF, DOCX, or TXT.")
    if not data:
        log_event("document_rejected", reason="empty_file")
        return DocumentResult(ok=False, filename=filename, error="File is empty.")
    if len(data) > MAX_BYTES:
        log_event("document_rejected", reason="too_large", bytes=len(data))
        return DocumentResult(ok=False, filename=filename, error="File exceeds the 5 MB limit.")

    try:
        if ext == ".pdf":
            text, doc_type = _extract_pdf(data), "pdf"
        elif ext == ".docx":
            text, doc_type = _extract_docx(data), "docx"
        else:
            text, doc_type = _extract_txt(data), "txt"
    except Exception as exc:  # corrupt/unreadable file
        log_event("extraction_failed", ext=ext, error_type=type(exc).__name__)
        return DocumentResult(ok=False, filename=filename, error=f"Could not read the {ext} file (it may be corrupt).")

    text = text.strip()
    if len(text) < MIN_CHARS:
        log_event("document_rejected", reason="no_text", chars=len(text))
        return DocumentResult(ok=False, filename=filename, doc_type=doc_type, error="No readable text found in the document.")

    spans = scan_for_injection(text)
    log_event(
        "document_accepted",
        doc_type=doc_type,
        chars=len(text),
        injection_flagged=bool(spans),
    )
    return DocumentResult(
        ok=True,
        text=text,
        doc_type=doc_type,
        filename=filename,
        char_count=len(text),
        injection_flagged=bool(spans),
        injection_spans=spans,
    )


def ingest_text(text: str, source_name: str = "pasted-text") -> DocumentResult:
    """Ingest pasted text (the paste-a-record path). Same guardrails as files."""
    text = (text or "").strip()
    if len(text) < MIN_CHARS:
        log_event("document_rejected", reason="no_text", chars=len(text))
        return DocumentResult(ok=False, filename=source_name, error="Please paste at least a few words of text.")
    spans = scan_for_injection(text)
    log_event("document_accepted", doc_type="text", chars=len(text), injection_flagged=bool(spans))
    return DocumentResult(
        ok=True,
        text=text,
        doc_type="text",
        filename=source_name,
        char_count=len(text),
        injection_flagged=bool(spans),
        injection_spans=spans,
    )


def ingest_path(path: str | Path) -> DocumentResult:
    """Convenience for CLI/eval: read a file from disk and ingest it."""
    p = Path(path)
    try:
        data = p.read_bytes()
    except OSError as exc:
        return DocumentResult(ok=False, filename=str(p), error=f"Could not open file: {exc}")
    return ingest_bytes(data, p.name)
